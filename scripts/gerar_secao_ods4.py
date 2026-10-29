from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('/home/ubuntu/palaciomental-current/docs')
OUT.mkdir(exist_ok=True)
DOCX = OUT / 'secao-7-ods-4-contexto-mestre.docx'
PURPLE = '321943'
GREEN = '1FAE6D'
BEIGE = 'F7F4EE'
GRAY = '5B5660'


def run_style(run, size=11, bold=False, color='000000', italic=False):
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def margins(cell, top=90, start=100, bottom=90, end=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for tag, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + tag))
        if node is None:
            node = OxmlElement('w:' + tag); tcMar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')


def border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ['top','left','bottom','right']:
        tag = qn('w:' + edge)
        el = borders.find(tag)
        if el is None:
            el = OxmlElement('w:' + edge); borders.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D6CBDC')


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text); run_style(r, size=14 if level == 1 else 12, bold=True, color=PURPLE)
    return p


def paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.line_spacing = 1.5; p.paragraph_format.space_after = Pt(after)
    run_style(p.add_run(text))
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; shade(c, PURPLE); margins(c); border(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing = 1.0
        run_style(p.add_run(h), size=9, bold=True, color='FFFFFF')
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            c = cells[i]; shade(c, BEIGE if ri % 2 == 0 else 'FFFFFF'); margins(c); border(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]; p.paragraph_format.line_spacing = 1.0
            run_style(p.add_run(value), size=9)
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27); sec.page_height = Inches(11.69)
sec.top_margin = Inches(1.18); sec.bottom_margin = Inches(0.79)
sec.left_margin = Inches(1.18); sec.right_margin = Inches(0.79)
normal = doc.styles['Normal']
normal.font.name = 'Arial'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); normal.font.size = Pt(11); normal.paragraph_format.line_spacing = 1.5

# Capa
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(90)
run_style(p.add_run('PALÁCIO MENTAL'), size=20, bold=True, color=PURPLE)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_style(p.add_run('SEÇÃO 7 — ODS E IMPACTO SOCIAL'), size=14, bold=True, color=GREEN)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(28)
run_style(p.add_run('Texto final para o Contexto Mestre'), size=12, color=PURPLE)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(120)
run_style(p.add_run('Versão para revisão da equipe'), size=11, color=GRAY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_style(p.add_run('Outubro de 2026'), size=11, color=GRAY)
doc.add_page_break()

heading(doc, '7. ODS — Objetivos de Desenvolvimento Sustentável', 1)
heading(doc, '7.1 ODS principal: ODS 4 — Educação de Qualidade', 2)
paragraph(doc, 'O Palácio Mental se relaciona principalmente à ODS 4 por apoiar experiências de aprendizagem baseadas em projetos. A plataforma ajuda pessoas e organizações a criar, contextualizar, desenvolver, documentar e preservar projetos para que o conhecimento produzido não se perca ao final de uma disciplina, programa ou ciclo de formação.')
paragraph(doc, 'Essa relação é real, mas indireta. O Palácio Mental não substitui escolas, universidades ou professores, não oferece certificação educacional e não garante aprendizagem por si só. Sua contribuição está em apoiar a organização da experiência, a colaboração, o registro dos aprendizados e a continuidade do conhecimento.')

heading(doc, '7.2 Conexão com as metas da ODS 4', 2)
table(doc, ['Meta', 'Relação com o Palácio Mental', 'Limite da contribuição'], [
    ['4.4 — Competências', 'Projetos permitem praticar, registrar decisões, desenvolver soluções, colaborar e apresentar evidências de competências técnicas, profissionais e empreendedoras.', 'A plataforma documenta e organiza a prática; não oferece, sozinha, formação profissional ou empregabilidade.'],
    ['4.7 — Desenvolvimento sustentável', 'Projetos de sustentabilidade, cidadania, direitos humanos, cultura, inclusão e desenvolvimento comunitário podem ser organizados, compartilhados e continuados.', 'O P.M. cria condições para esses projetos; não garante que todo projeto produzirá impacto social.'],
], widths=[1.35,3.45,2.2])

heading(doc, '7.3 Forma de contribuição', 2)
paragraph(doc, 'A contribuição esperada ocorre por meio do seguinte ciclo:')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(10)
run_style(p.add_run('projeto → prática → registro → reflexão → colaboração → continuidade'), size=12, bold=True, color=GREEN)
paragraph(doc, 'No ambiente institucional, esse ciclo pode apoiar disciplinas, programas, grupos, laboratórios, projetos de extensão e iniciativas de organizações sociais. A autoria permanece com as pessoas e os projetos institucionais permanecem separados da vitrine global, conforme a arquitetura do produto.')

heading(doc, '7.4 Indicadores iniciais de impacto', 2)
paragraph(doc, 'A relação com a ODS 4 ainda é uma hipótese de impacto. Para avaliá-la, o Palácio Mental poderá acompanhar:')
for item in [
    'projetos educacionais criados e contextualizados;',
    'participantes, professores e mentores envolvidos;',
    'colaborações e interações qualificadas;',
    'projetos reutilizados, continuados ou descobertos por outras pessoas;',
    'projetos relacionados a sustentabilidade, inclusão ou desenvolvimento comunitário;',
    'percepção dos participantes sobre organização, aprendizado e continuidade.',
]:
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent = Inches(0.35); p.paragraph_format.line_spacing = 1.5; p.paragraph_format.space_after = Pt(2); run_style(p.add_run(item))
paragraph(doc, 'Esses indicadores não medem sozinhos a qualidade da educação. Eles mostram se a plataforma está contribuindo para práticas de aprendizagem, colaboração e preservação do conhecimento.')

heading(doc, '7.5 Síntese para comunicação', 2)
paragraph(doc, 'O Palácio Mental contribui para a ODS 4 ao apoiar a aprendizagem baseada em projetos, o desenvolvimento de competências e a continuidade do conhecimento. Essa contribuição será validada por evidências de uso, participação, contextualização, colaboração e reaproveitamento de projetos.')

heading(doc, '7.6 Referência oficial', 2)
paragraph(doc, 'A ODS 4 busca assegurar educação inclusiva, equitativa e de qualidade e promover oportunidades de aprendizagem ao longo da vida para todas e todos. As metas 4.4 e 4.7 são as conexões prioritárias desta seção.')
paragraph(doc, 'Fonte: ONU Brasil. Objetivo de Desenvolvimento Sustentável 4 — Educação de qualidade. Disponível em: https://brasil.un.org/pt-br/sdgs/4. Acesso em: 27 out. 2026.')

for section in doc.sections:
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_style(footer.add_run('Palácio Mental — seção de revisão do Contexto Mestre'), size=8, color=GRAY)

doc.save(DOCX)
print(DOCX)
